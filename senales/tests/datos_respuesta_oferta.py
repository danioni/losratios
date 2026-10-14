"""Un mundo sintético y determinista para los tests de la fase N1. Sin red.

Las fuentes se fabrican con la forma que describen FUENTES.md (N0.16) y los
lectores de fuentes_respuesta_oferta.py: el CSV sin encabezado de la FHFA, el
xlsx de terminadas del Censo con sus hojas, la copia del BLS en la grilla del
"SeriesReport", el CSV de la B.103 del Z.1 y su HTML, y los insumos publicados
de las fases R y N0. Los valores no son los reales: son coherentes entre sí
para que los gates cierren por construcción, salvo donde el test quiere que
no cierren.
"""

from __future__ import annotations

import math
import random
from pathlib import Path

import pandas as pd

from senales.configuracion import ONZAS_TROY_POR_TONELADA

SEMILLA = 20261014
ANIO_INICIO_PINK = 1960
ANIO_FIN = 2025
ANIO_FIN_DS140 = {"oro": 2022, "plata": 2021}
MESES_PINK_EN_DISPUTA = ("1985-03",)
ANIOS_DS140_FUERA = (1968, 1969)  # el control contra el Pink Sheet los marca
ANIOS_PLATA_EN_DISPUTA = (2020, 2021, 2022, 2024)  # el control del BGS de N0


def _ruido(generador: random.Random, sigma: float) -> float:
    return generador.gauss(0.0, sigma)


def pink_sheet_mensual(generador: random.Random, inicial: float) -> dict[str, float]:
    valores, nivel = {}, math.log(inicial)
    for anio in range(ANIO_INICIO_PINK, ANIO_FIN + 2):
        for mes in range(1, 13):
            if anio == ANIO_FIN + 1 and mes > 9:
                break
            nivel += 0.004 + _ruido(generador, 0.03)
            valores[f"{anio}-{mes:02d}"] = round(math.exp(nivel), 2)
    return valores


def cpi_mensual(generador: random.Random) -> dict[str, float]:
    valores, nivel = {}, math.log(9.8)
    for anio in range(1913, ANIO_FIN + 2):
        for mes in range(1, 13):
            if anio == ANIO_FIN + 1 and mes > 8:
                break
            nivel += 0.0025 + _ruido(generador, 0.002)
            valores[f"{anio}-{mes:02d}"] = round(math.exp(nivel), 3)
    return valores


def promedio(valores: dict[str, float], anio: int, periodos: int = 12) -> float | None:
    del_anio = [v for k, v in valores.items() if k.startswith(f"{anio}-")]
    return sum(del_anio) / periodos if len(del_anio) == periodos else None


def produccion(generador: random.Random, precio_real_anual: dict[int, float], inicial: float, desde: int = 1900) -> dict[int, float]:
    """Un flujo que responde con rezago al precio real (β_1 = 0.3) más ruido."""
    valores, nivel = {}, math.log(inicial)
    for anio in range(desde, ANIO_FIN + 1):
        previo, anterior = precio_real_anual.get(anio - 1), precio_real_anual.get(anio - 2)
        respuesta = 0.3 * (math.log(previo) - math.log(anterior)) if previo and anterior else 0.0
        nivel += 0.01 + respuesta + _ruido(generador, 0.02)
        valores[anio] = round(math.exp(nivel), 0)
    return valores


def mundo(fecha: str = "2026-10-14") -> dict:
    generador = random.Random(SEMILLA)
    cpi = cpi_mensual(generador)
    cpi_anual = {a: promedio(cpi, a) for a in range(1913, ANIO_FIN + 1)}
    pink = {"oro": pink_sheet_mensual(generador, 35.0), "plata": pink_sheet_mensual(generador, 0.91)}
    pink_anual = {m: {a: promedio(v, a) for a in range(ANIO_INICIO_PINK, ANIO_FIN + 1)} for m, v in pink.items()}
    ds140 = {}
    for metal, inicial in (("oro", 609000.0), ("plata", 20000.0)):
        valores, nivel = {}, math.log(inicial / ONZAS_TROY_POR_TONELADA)
        for anio in range(1900, ANIO_FIN_DS140[metal] + 1):
            if anio >= ANIO_INICIO_PINK:
                por_oz = pink_anual[metal][anio] * (1.10 if anio in ANIOS_DS140_FUERA else 1.004)
            else:
                nivel += 0.003 + _ruido(generador, 0.01)
                por_oz = math.exp(nivel)
            valores[anio] = round(por_oz * ONZAS_TROY_POR_TONELADA, 0)
        ds140[metal] = valores
    real = {m: {a: (pink_anual[m][a] / cpi_anual[a]) for a in pink_anual[m] if pink_anual[m][a] and cpi_anual.get(a)} for m in pink}
    prod = {"oro": produccion(generador, real["oro"], 386.0), "plata": produccion(generador, real["plata"], 5400.0)}
    fhfa = {}
    nivel = math.log(60.04)
    for anio in range(1975, ANIO_FIN + 2):
        for trimestre in range(1, 5):
            if anio == ANIO_FIN + 1 and trimestre > 2:
                break
            nivel += 0.012 + _ruido(generador, 0.01)
            fhfa[f"{anio}:Q{trimestre}"] = round(math.exp(nivel), 2)
    fhfa_anual = {a: promedio({k.replace(":Q", "-"): v for k, v in fhfa.items()}, a, 4) for a in range(1975, ANIO_FIN + 1)}
    real_fhfa = {a: fhfa_anual[a] / cpi_anual[a] for a in fhfa_anual if fhfa_anual[a]}
    terminadas_anual = produccion(generador, real_fhfa, 1300.0, desde=1968)
    terminadas_mensual = {}
    for anio, total in terminadas_anual.items():
        base = round(total / 12.0, 1)
        meses = [base] * 12
        meses[11] = round(total - sum(meses[:11]), 1)
        terminadas_mensual[anio] = meses
    b103 = {}
    acciones, fdi, patrimonio = 100000.0, 20000.0, 60000.0
    for anio in range(1945, ANIO_FIN + 2):
        for trimestre in range(1, 5):
            if anio == 1945 and trimestre < 4:
                continue
            if anio == ANIO_FIN + 1 and trimestre > 2:
                break
            acciones *= math.exp(0.02 + _ruido(generador, 0.04))
            fdi *= math.exp(0.02 + _ruido(generador, 0.02))
            patrimonio *= math.exp(0.015 + _ruido(generador, 0.01))
            b103[f"{anio}:Q{trimestre}"] = (round(acciones, 1), round(fdi, 1), round(patrimonio, 1), round((round(acciones, 1) + round(fdi, 1)) / round(patrimonio, 1) * 100.0, 2))
    valuacion_anual = {a: sum(b103[f"{a}:Q{q}"][0] / b103[f"{a}:Q{q}"][2] * 100.0 for q in range(1, 5)) / 4.0 for a in range(1946, ANIO_FIN + 1)}
    tasa = {}
    for anio in range(1946, ANIO_FIN + 1):
        previo, anterior = valuacion_anual.get(anio - 1), valuacion_anual.get(anio - 2)
        respuesta = 0.3 * (math.log(previo) - math.log(anterior)) if previo and anterior else 0.0
        tasa[anio] = round(0.5 + 100 * respuesta + _ruido(generador, 0.8), 6)
    return {
        "fecha": fecha, "cpi": cpi, "pink": pink, "ds140": ds140, "produccion": prod, "fhfa": fhfa,
        "terminadas_anual": terminadas_anual, "terminadas_mensual": terminadas_mensual, "b103": b103, "tasa": tasa,
    }


# --- Escritura de los archivos con la forma de cada fuente ----------------------------------


def escribir_fhfa(ruta: Path, fhfa: dict[str, float]) -> None:
    lineas = [f"USA,{p[:4]},{p[-1]},{v:.2f}" for p, v in fhfa.items()]
    lineas += [f"DV_ENC,{p[:4]},{p[-1]},{v * 0.9:.2f}" for p, v in fhfa.items()]  # otra área, que el lector ignora
    ruta.write_text("\n".join(lineas) + "\n", encoding="utf-8")


def escribir_fred(ruta: Path, id_serie: str, valores: dict[str, float], decimales: int) -> None:
    lineas = [f"observation_date,{id_serie}"]
    for periodo, valor in valores.items():
        fecha = periodo if len(periodo) == 10 else (f"{periodo[:4]}-{(int(periodo[-1]) - 1) * 3 + 1:02d}-01" if ":Q" in periodo else f"{periodo}-01")
        lineas.append(f"{fecha},{valor:.{decimales}f}")
    ruta.write_text("\n".join(lineas) + "\n", encoding="utf-8")


def escribir_terminadas(ruta: Path, anual: dict[int, float], mensual: dict[int, list[float]]) -> None:
    import openpyxl

    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "Annual"
    hoja.append(["New Privately-Owned Housing Units Completed", None, None])
    hoja.append(["(Thousands of units)", None, None])
    hoja.append(["Year", "Total", "1 unit", "5 units or more"])
    for anio, total in anual.items():
        hoja.append([anio, total, round(total * 0.7, 1), round(total * 0.3, 1)])
    hoja.append(["Source: U.S. Census Bureau and U.S. Department of Housing and Urban Development"])
    hoja_mensual = libro.create_sheet("Monthly Not Seasonally Adjusted")
    hoja_mensual.append(["(Thousands of units)"])
    hoja_mensual.append(["Year", "Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"])
    for anio, meses in mensual.items():
        hoja_mensual.append([anio, *meses])
    hoja_saar = libro.create_sheet("Monthly SAAR")
    hoja_saar.append(["Year", "Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"])
    for anio, meses in mensual.items():
        hoja_saar.append([anio, *[round(m * 12, 0) for m in meses]])
    libro.save(ruta)


def escribir_bls(ruta: Path, cpi: dict[str, float], con_anual: bool = True) -> None:
    import openpyxl

    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "BLS Data Series"
    hoja.append(["Consumer Price Index for All Urban Consumers (CPI-U)"])
    hoja.append(["Series Id:", "CUUR0000SA0"])
    hoja.append(["Not Seasonally Adjusted"])
    hoja.append(["Year", "Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"] + (["Annual", "HALF1", "HALF2"] if con_anual else []))
    anios = sorted({int(k[:4]) for k in cpi})
    for anio in anios:
        meses = [cpi.get(f"{anio}-{m:02d}") for m in range(1, 13)]
        fila = [anio, *meses]
        if con_anual:
            completos = [m for m in meses if m is not None]
            fila += [round(sum(completos) / 12.0, 3) if len(completos) == 12 else None, None, None]
        hoja.append(fila)
    libro.save(ruta)


def escribir_b103_csv(ruta: Path, b103: dict[str, tuple]) -> None:
    lineas = ["date,LM103164105.Q,LM103192105.Q,FL102090005.Q,FL103164106.Q"]
    for periodo, (acciones, fdi, patrimonio, razon) in b103.items():
        lineas.append(f"{periodo},{acciones:.1f},{fdi:.1f},{patrimonio:.1f},{razon:.2f}")
    ruta.write_text("\n".join(lineas) + "\n", encoding="utf-8")


def escribir_b103_html(ruta: Path, b103: dict[str, tuple]) -> None:
    periodos = [p for p in b103 if p.startswith(("2025:", "2026:"))][-6:]
    anuales = ["2024", "2025"]
    encabezado = "".join(f"<th>{p}</th>" for p in anuales + periodos)
    filas = []
    for linea, (mnemonico, posicion, nombre) in enumerate((("LM103164105", 0, "Corporate equities"), ("LM103192105", 1, "Foreign direct investment"), ("FL102090005", 2, "Net worth"), ("FL103164106", 3, "Equities as a percentage of net worth")), start=1):
        celdas = []
        for periodo in anuales + periodos:
            clave = f"{periodo}:Q4" if ":Q" not in periodo else periodo
            valor = b103[clave][posicion]
            celdas.append(f"{valor / 1000.0:,.1f}" if posicion < 3 else f"{valor:.1f}")
        filas.append(f"<tr><td>{linea}</td><td>{nombre}</td><td>{mnemonico}</td>" + "".join(f"<td>{c}</td>" for c in celdas) + "</tr>")
    ruta.write_text(f"<html><body><table><tr><th>Line</th><th></th><th></th>{encabezado}</tr>{''.join(filas)}</table></body></html>", encoding="utf-8")


def escribir_ds140(ruta: Path, metal: str, ds140: dict[int, float], produccion: dict[int, float]) -> None:
    import openpyxl

    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = {"oro": "Gold", "plata": "Silver"}[metal]
    hoja.append([f"{'GOLD' if metal == 'oro' else 'SILVER'} STATISTICS1"])
    hoja.append(["U.S. GEOLOGICAL SURVEY"])
    hoja.append(["[All values are in metric tons (t) gold content unless otherwise noted]"])
    hoja.append(["Last modification: test"])
    hoja.append(["Year", "Primary production", "Unit value ($/t)", "Unit value (98$/t)", "World production"])
    for anio, valor in ds140.items():
        hoja.append([anio, 100, valor, valor * 0.5, produccion.get(anio, "NA")])
    hoja.append(["NA Not available."])
    libro.save(ruta)


def escribir_precios_mensuales(ruta: Path, pink: dict[str, dict[str, float]]) -> None:
    filas = []
    for mes in pink["oro"]:
        disputa = mes in MESES_PINK_EN_DISPUTA
        filas.append(
            {
                "mes": mes, "oro_usd_oz": pink["oro"][mes], "oro_error_redondeo_pct": 0.01, "oro_definicion": "fixing de la tarde de Londres",
                "oro_contraste_fmi": "valor en disputa: Pink Sheet x, FMI y" if disputa else ("sin comparar" if mes < "1980-01" else "dentro del umbral"),
                "plata_usd_oz": pink["plata"][mes], "plata_error_redondeo_pct": 0.05, "plata_estado": "estimación",
                "plata_contraste_fmi": "valor en disputa: Pink Sheet x, FMI y" if disputa else ("sin comparar" if mes < "1980-01" else "dentro del umbral"),
                "pink_sheet_edicion": "2025-01-03" if mes < "2025-01" else "vigente", "btc_usd": "",
            }
        )
    pd.DataFrame(filas).to_csv(ruta, index=False, lineterminator="\n")


def escribir_numerador_series(ruta: Path, produccion: dict[str, dict[int, float]], tasa: dict[int, float]) -> None:
    filas = []
    for metal, valores in produccion.items():
        for anio, valor in valores.items():
            filas.append(
                {
                    "serie": f"{metal}_produccion_mundial_t", "anio": anio, "fecha": "", "valor": valor, "unidad": "toneladas",
                    "estado": "estimación" if anio == ANIO_FIN else "dato",
                    "control": "valor en disputa" if (metal == "plata" and anio in ANIOS_PLATA_EN_DISPUTA) else ("dentro del umbral" if anio >= 2020 else ""),
                    "cita": "USGS (sintético)", "nota": "",
                }
            )
    for anio, valor in tasa.items():
        filas.append({"serie": "acciones_eeuu_emision_neta_no_financieras_pct_vm", "anio": anio, "fecha": "", "valor": valor, "unidad": "% del valor de mercado previo", "estado": "dato", "control": "", "cita": "Z.1 (sintético)", "nota": ""})
    pd.DataFrame(filas).to_csv(ruta, index=False, lineterminator="\n")


def escribir_mundo(raiz: Path, m: dict) -> dict[str, Path]:
    """Escribe todos los crudos y los insumos publicados del mundo sintético bajo `raiz`."""
    crudo, privado, series = raiz / "data" / "raw", raiz / "data" / "privado" / "raw", raiz / "data" / "series"
    for directorio in (crudo, privado, series):
        directorio.mkdir(parents=True, exist_ok=True)
    fecha = m["fecha"]
    escribir_fhfa(crudo / f"fhfa_hpi_at_eeuu_{fecha}.csv", m["fhfa"])
    escribir_fred(privado / f"fred_ussthpi_{fecha}.csv", "USSTHPI", m["fhfa"], 2)
    escribir_fred(crudo / f"fred_cpiaucns_{fecha}.csv", "CPIAUCNS", m["cpi"], 3)
    escribir_bls(privado / f"bls_cpi_u_{fecha}.xlsx", m["cpi"])
    escribir_terminadas(crudo / f"censo_viviendas_terminadas_{fecha}.xlsx", m["terminadas_anual"], m["terminadas_mensual"])
    escribir_b103_csv(crudo / f"z1_b103_{fecha}.csv", m["b103"])
    escribir_b103_html(privado / f"z1_html_b103_{fecha}.htm", m["b103"])
    for metal in ("oro", "plata"):
        escribir_ds140(crudo / f"usgs_ds140_{metal}_{fecha}.xlsx", metal, m["ds140"][metal], m["produccion"][metal])
    escribir_precios_mensuales(series / "precios_mensuales.csv", m["pink"])
    escribir_numerador_series(series / "numerador_series.csv", m["produccion"], m["tasa"])
    return {"crudo": crudo, "privado": privado, "series": series}
