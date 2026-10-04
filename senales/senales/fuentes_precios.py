"""Descarga y lectura de las fuentes de precios de la fase R (ratios).

Tres reglas que este módulo no negocia:

1. El crudo se guarda en data/raw/, como los de S2, si la licencia de la fuente
   permite redistribuirlo (CC BY, CC BY-NC). Si no, va a un directorio ignorado
   por git: el repositorio no lo redistribuye (A-R0-15).
2. De cada descarga queda la URL, la fecha y el SHA-256 en un manifiesto que sí
   se publica. Si el crudo que hay en disco no coincide con el hash publicado,
   la corrida se detiene.
3. Una fuente de contraste se lee, se compara y se descarta. Las funciones
   `contraste_*` devuelven los datos en memoria y no escriben nada (A-R0-12).
"""

from __future__ import annotations

import hashlib
import json
import re
import time
from dataclasses import dataclass
from datetime import date, datetime, timedelta, timezone
from email.utils import parsedate_to_datetime
from io import StringIO
from pathlib import Path
from urllib.parse import urljoin

import pandas as pd
import requests

from senales.configuracion import (
    COLUMNAS_DESCARGAS,
    FMI_CODIGO_ORO,
    FMI_CODIGO_PLATA,
    FMI_DESCRIPCION_ORO,
    FMI_DESCRIPCION_PLATA,
    FMI_HOJA,
    PINK_SHEET_DESCRIPCION_ORO,
    PINK_SHEET_DESCRIPCION_PLATA,
    PINK_SHEET_HOJA_DESCRIPCION,
    PINK_SHEET_HOJA_PRECIOS,
    PINK_SHEET_UNIDAD,
    Contraste,
    CopiaManual,
    Descarga,
    EdicionCongelada,
)
from senales.fuentes_fred import ErrorDeFuente
from senales.nucleo import normalizar_texto

TIEMPO_LIMITE = 90

# Segundos entre dos pedidos seguidos al mismo host. Las fuentes son gratuitas y
# nadie les debe velocidad a este script.
PAUSA_ENTRE_PEDIDOS = 2.0

# api.nasdaq.com reinicia la conexión si el pedido no trae cabeceras de
# navegador (FUENTES.md, sección 3.2). Solo se usan para ese host.
CABECERAS_NAVEGADOR = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/128.0 Safari/537.36"
    ),
    "Accept": "application/json, text/plain, */*",
    "Accept-Language": "en-US,en;q=0.9",
}

# Primeros bytes de cada formato. Una página de error con HTTP 200 no pasa.
FIRMAS = {
    "xlsx": (b"PK",),
    "xls": (b"\xd0\xcf\x11\xe0",),
    "json": (b"{",),
}

MESES_INGLES = {
    nombre: numero
    for numero, nombre in enumerate(
        [
            "january", "february", "march", "april", "may", "june", "july",
            "august", "september", "october", "november", "december",
        ],
        start=1,
    )
}

_ultimo_pedido: dict[str, float] = {}


@dataclass(frozen=True)
class RegistroDescarga:
    """Lo que se sabe de una descarga. Es una fila del manifiesto."""

    descarga: Descarga
    fecha_descarga: date
    ruta: Path
    url: str
    bytes: int
    sha256: str
    actualizada: date | None  # la fecha que la fuente declara, si declara una
    descargada_ahora: bool

    def fila(self) -> dict[str, object]:
        return {
            "fecha_descarga": self.fecha_descarga.isoformat(),
            "fuente": self.descarga.clave,
            "clase_licencia": self.descarga.clase_licencia,
            "licencia": self.descarga.licencia,
            "url": self.url,
            "bytes": self.bytes,
            "sha256": self.sha256,
            "actualizada": "" if self.actualizada is None else self.actualizada.isoformat(),
            "crudo_en_repo": "sí" if self.descarga.crudo_versionado else "no",
        }

    def linea(self) -> str:
        accion = "descargada" if self.descargada_ahora else "reutilizada"
        if self.fecha_descarga != date.today() and not self.descargada_ahora:
            accion = f"copia del {self.fecha_descarga}"
        lugar = "data/raw" if self.descarga.crudo_versionado else "fuera del repositorio"
        return (
            f"{self.descarga.clave}: {accion} {self.ruta.name} ({lugar}), "
            f"{self.bytes} bytes, sha256 {self.sha256[:16]}…"
        )


# --- Pedidos -----------------------------------------------------------------


def _pedir(url: str, sesion=None, cabeceras: dict | None = None, parametros: dict | None = None):
    """Un GET con pausa entre pedidos al mismo host. Falla con ErrorDeFuente."""
    host = url.split("/")[2]
    espera = PAUSA_ENTRE_PEDIDOS - (time.monotonic() - _ultimo_pedido.get(host, 0.0))
    if host in _ultimo_pedido and espera > 0:
        time.sleep(espera)
    cliente = sesion or requests
    try:
        respuesta = cliente.get(url, timeout=TIEMPO_LIMITE, headers=cabeceras, params=parametros)
        respuesta.raise_for_status()
    except requests.RequestException as error:
        raise ErrorDeFuente(f"no se pudo leer {url}: {error}") from error
    finally:
        _ultimo_pedido[host] = time.monotonic()
    return respuesta


def sha256_de(ruta: Path) -> str:
    resumen = hashlib.sha256()
    with ruta.open("rb") as archivo:
        for bloque in iter(lambda: archivo.read(1 << 20), b""):
            resumen.update(bloque)
    return resumen.hexdigest()


# --- Descarga y manifiesto ---------------------------------------------------


def ruta_cruda(descarga: Descarga, fecha_descarga: date, dir_abierto: Path, dir_privado: Path) -> Path:
    """A-R0-15: el directorio del crudo lo decide la clase de licencia."""
    directorio = dir_abierto if descarga.crudo_versionado else dir_privado
    return directorio / f"{descarga.clave}_{fecha_descarga.isoformat()}.{descarga.extension}"


def extraer_enlace(html: str, descarga: Descarga) -> str:
    """Busca en la página de la fuente el enlace al archivo."""
    coincidencia = re.search(descarga.patron_enlace or "", html)
    if descarga.patron_enlace is None or coincidencia is None:
        raise ErrorDeFuente(
            f"{descarga.clave}: la página {descarga.url} no tiene un enlace que "
            f"coincida con {descarga.patron_enlace!r}. La fuente cambió de lugar; "
            "revisar FUENTES.md antes de apuntar a otra URL."
        )
    return urljoin(descarga.url, coincidencia.group(1).replace("&amp;", "&"))


def _fecha_de_cabecera(valor: str | None) -> date | None:
    if not valor:
        return None
    try:
        return parsedate_to_datetime(valor).date()
    except (TypeError, ValueError):
        return None


def _verificar_formato(descarga: Descarga, contenido: bytes, url: str) -> None:
    if not contenido.strip():
        raise ErrorDeFuente(f"{descarga.clave}: {url} devolvió una respuesta vacía")
    firmas = FIRMAS.get(descarga.extension)
    if firmas is not None and not contenido.lstrip().startswith(firmas):
        raise ErrorDeFuente(
            f"{descarga.clave}: {url} no devolvió un archivo .{descarga.extension} "
            f"(empieza con {contenido[:12]!r})"
        )
    if descarga.extension == "csv" and b"," not in contenido.splitlines()[0]:
        raise ErrorDeFuente(f"{descarga.clave}: {url} no devolvió un CSV")


def leer_manifiesto(ruta: Path) -> pd.DataFrame:
    if not ruta.exists():
        return pd.DataFrame(columns=COLUMNAS_DESCARGAS)
    return pd.read_csv(ruta, dtype=str, keep_default_na=False)


def _fila_del_manifiesto(manifiesto: pd.DataFrame, descarga: Descarga, fecha: date):
    filas = manifiesto.loc[
        (manifiesto["fecha_descarga"] == fecha.isoformat())
        & (manifiesto["fuente"] == descarga.clave)
    ]
    return None if filas.empty else filas.iloc[-1]


def descargar(
    descarga: Descarga,
    fecha_descarga: date,
    dir_abierto: Path,
    dir_privado: Path,
    manifiesto: pd.DataFrame,
    sesion=None,
) -> RegistroDescarga:
    """Baja el crudo de una fuente, o reutiliza el que ya está.

    El crudo de un día no se pisa. Si ya existe se reutiliza, y si el manifiesto
    tiene su fila, el hash tiene que coincidir: un crudo que cambió debajo de un
    hash publicado no es el archivo del que salió la serie.
    """
    destino = ruta_cruda(descarga, fecha_descarga, dir_abierto, dir_privado)
    fila = _fila_del_manifiesto(manifiesto, descarga, fecha_descarga)

    if destino.exists():
        resumen = sha256_de(destino)
        if fila is not None and fila["sha256"] != resumen:
            raise ErrorDeFuente(
                f"{descarga.clave}: {destino.name} tiene sha256 {resumen[:16]}… y el "
                f"manifiesto publicó {fila['sha256'][:16]}… para esa descarga. El crudo "
                "cambió; no se usa."
            )
        actualizada = None
        if fila is not None and fila["actualizada"]:
            actualizada = date.fromisoformat(fila["actualizada"])
        return RegistroDescarga(
            descarga=descarga,
            fecha_descarga=fecha_descarga,
            ruta=destino,
            url=descarga.url if fila is None else fila["url"],
            bytes=destino.stat().st_size,
            sha256=resumen,
            actualizada=actualizada,
            descargada_ahora=False,
        )

    if fila is not None and not descarga.crudo_versionado:
        raise ErrorDeFuente(
            f"{descarga.clave}: la descarga del {fecha_descarga} está en el manifiesto "
            f"pero su crudo no está en {destino.parent}. Los crudos que no se pueden "
            f"redistribuir no viajan con el repositorio (A-R0-15): hay que bajarlo de "
            f"{fila['url']} y comprobar que su sha256 sea {fila['sha256']}."
        )

    url = descarga.url
    if descarga.patron_enlace is not None:
        url = extraer_enlace(_pedir(descarga.url, sesion).text, descarga)
    respuesta = _pedir(url, sesion)
    contenido = respuesta.content
    _verificar_formato(descarga, contenido, url)

    destino.parent.mkdir(parents=True, exist_ok=True)
    temporal = destino.with_suffix(destino.suffix + ".tmp")
    temporal.write_bytes(contenido)
    temporal.replace(destino)
    return RegistroDescarga(
        descarga=descarga,
        fecha_descarga=fecha_descarga,
        ruta=destino,
        url=url,
        bytes=len(contenido),
        sha256=hashlib.sha256(contenido).hexdigest(),
        actualizada=_fecha_de_cabecera(respuesta.headers.get("Last-Modified")),
        descargada_ahora=True,
    )


def cargar_edicion_congelada(
    edicion: EdicionCongelada, dir_crudo: Path, sesion=None
) -> RegistroDescarga:
    """La copia versionada de una edición congelada, verificada contra su hash.

    El pipeline no depende de que la URL de la edición siga en línea: usa la
    copia que viaja con el repositorio (A-R0-19). Solo si esa copia falta sale a
    buscarla, y la acepta únicamente si su SHA-256 es el esperado.
    """
    ruta = dir_crudo / edicion.archivo
    descargada = False
    if not ruta.exists():
        respuesta = _pedir(edicion.descarga.url, sesion)
        contenido = respuesta.content
        _verificar_formato(edicion.descarga, contenido, edicion.descarga.url)
        if hashlib.sha256(contenido).hexdigest() != edicion.sha256:
            raise ErrorDeFuente(
                f"{edicion.descarga.clave}: {edicion.descarga.url} ya no entrega la edición "
                f"del {edicion.fecha_edicion} (el sha256 no es {edicion.sha256[:16]}…). Sin la "
                f"copia versionada de {edicion.archivo} no se puede reconstruir el empalme."
            )
        ruta.parent.mkdir(parents=True, exist_ok=True)
        temporal = ruta.with_suffix(ruta.suffix + ".tmp")
        temporal.write_bytes(contenido)
        temporal.replace(ruta)
        descargada = True

    resumen = sha256_de(ruta)
    if resumen != edicion.sha256:
        raise ErrorDeFuente(
            f"{edicion.descarga.clave}: {ruta.name} tiene sha256 {resumen[:16]}… y la edición "
            f"congelada es {edicion.sha256[:16]}…. El archivo cambió; no se usa."
        )
    return RegistroDescarga(
        descarga=edicion.descarga,
        fecha_descarga=edicion.fecha_descarga,
        ruta=ruta,
        url=edicion.descarga.url,
        bytes=ruta.stat().st_size,
        sha256=resumen,
        actualizada=edicion.fecha_edicion,
        descargada_ahora=descargada,
    )


def cargar_copia_manual(copia: CopiaManual, dir_crudo: Path) -> RegistroDescarga:
    """La copia versionada de un archivo que se bajó una vez, a mano (A-R0-20).

    Nunca sale a la red. Si la copia falta o cambió, la corrida se detiene: el
    control que depende de ella decide qué meses entran a las métricas, y
    publicar sin él sería publicar otra cosa.
    """
    ruta = dir_crudo / copia.archivo
    if not ruta.exists():
        raise ErrorDeFuente(
            f"{copia.descarga.clave}: falta {copia.archivo} en {dir_crudo}. Este archivo no "
            "se baja solo: los términos de la fuente prohíben la descarga masiva por "
            f"medios automatizados (A-R0-20). Hay que bajarlo a mano de {copia.descarga.url} "
            f"y comprobar que su sha256 sea {copia.sha256}; si la fuente ya publica otra "
            "edición, actualizar configuracion.py con la fecha y el hash nuevos."
        )
    resumen = sha256_de(ruta)
    if resumen != copia.sha256:
        raise ErrorDeFuente(
            f"{copia.descarga.clave}: {ruta.name} tiene sha256 {resumen[:16]}… y la copia "
            f"declarada es {copia.sha256[:16]}…. El archivo cambió; no se usa."
        )
    return RegistroDescarga(
        descarga=copia.descarga,
        fecha_descarga=copia.fecha_descarga,
        ruta=ruta,
        url=copia.descarga.url,
        bytes=ruta.stat().st_size,
        sha256=resumen,
        actualizada=None,
        descargada_ahora=False,
    )


def actualizar_manifiesto(ruta: Path, registros: list[RegistroDescarga]) -> bool:
    """Deja una fila por (fecha, fuente). Devuelve True si el archivo cambió."""
    previo = leer_manifiesto(ruta)
    nuevas = pd.DataFrame([r.fila() for r in registros], columns=COLUMNAS_DESCARGAS).astype(str)
    claves = set(zip(nuevas["fecha_descarga"], nuevas["fuente"]))
    conservadas = previo.loc[
        [par not in claves for par in zip(previo["fecha_descarga"], previo["fuente"])]
    ]
    tabla = pd.concat([conservadas, nuevas], ignore_index=True)
    tabla = tabla.sort_values(["fecha_descarga", "fuente"], kind="stable")
    texto = tabla.loc[:, COLUMNAS_DESCARGAS].to_csv(index=False, lineterminator="\n")
    if ruta.exists() and ruta.read_text(encoding="utf-8") == texto:
        return False
    ruta.parent.mkdir(parents=True, exist_ok=True)
    temporal = ruta.with_suffix(ruta.suffix + ".tmp")
    temporal.write_text(texto, encoding="utf-8", newline="")
    temporal.replace(ruta)
    return True


# --- Utilidades de fechas ----------------------------------------------------


def a_mes(valor) -> pd.Timestamp:
    """El primer día del mes, que es como se indexa toda serie mensual."""
    return pd.Timestamp(valor).to_period("M").to_timestamp()


def _serie_mensual(pares: list[tuple[pd.Timestamp, float]], nombre: str) -> pd.Series:
    indice = pd.DatetimeIndex([mes for mes, _ in pares], name="mes")
    serie = pd.Series([valor for _, valor in pares], index=indice, name=nombre, dtype="float64")
    if serie.index.has_duplicates:
        raise ErrorDeFuente(f"{nombre}: la fuente trae meses repetidos")
    return serie.sort_index()


def _serie_diaria(fechas, valores, nombre: str) -> pd.Series:
    indice = pd.DatetimeIndex(pd.to_datetime(list(fechas)), name="fecha")
    serie = pd.Series(list(valores), index=indice, name=nombre, dtype="float64").dropna()
    serie = serie[~serie.index.duplicated(keep="last")].sort_index()
    if serie.empty:
        raise ErrorDeFuente(f"{nombre}: la serie quedó vacía después de limpiar")
    return serie


# --- Banco Mundial, Pink Sheet ----------------------------------------------


@dataclass(frozen=True)
class PinkSheet:
    oro: pd.Series
    plata: pd.Series
    actualizada: date | None


def filas_de_xlsx(ruta: Path, hoja: str) -> list[tuple]:
    import openpyxl

    libro = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    try:
        if hoja not in libro.sheetnames:
            raise ErrorDeFuente(f"{ruta.name}: no tiene la hoja '{hoja}' (tiene {libro.sheetnames})")
        return [tuple(fila) for fila in libro[hoja].iter_rows(values_only=True)]
    finally:
        libro.close()


def _fecha_de_actualizacion(filas: list[tuple]) -> date | None:
    """Lee 'Updated on October 02, 2026' del encabezado de la hoja de precios."""
    for fila in filas[:10]:
        for celda in fila:
            if not isinstance(celda, str):
                continue
            hallado = re.search(r"Updated on (\w+) (\d{1,2}), (\d{4})", celda)
            if hallado and hallado.group(1).lower() in MESES_INGLES:
                return date(
                    int(hallado.group(3)), MESES_INGLES[hallado.group(1).lower()], int(hallado.group(2))
                )
    return None


def _verificar_descripcion(filas: list[tuple], esperada: str, que: str) -> None:
    """La convención del Pink Sheet es lo que su hoja Description dice.

    Si la frase cambia, cambió lo que la serie mide, y la corrida se detiene.
    """
    objetivo = normalizar_texto(esperada)
    inicio = objetivo.split(",")[0]
    candidatas = [
        celda
        for fila in filas
        for celda in fila
        if isinstance(celda, str) and normalizar_texto(celda).startswith(inicio)
    ]
    if any(normalizar_texto(celda) == objetivo for celda in candidatas):
        return
    hallada = candidatas[0] if candidatas else "(no hay ninguna descripción que empiece así)"
    raise ErrorDeFuente(
        f"Pink Sheet: la descripción {que} ya no es la que se leyó al fijar la "
        f"convención.\n  esperada: {esperada}\n  hallada:  {hallada}\n"
        "No publicar con la definición vieja: releer la fuente y actualizar "
        "configuracion.py, FUENTES.md y SUPUESTOS.md (A-R0-7, A-R0-8)."
    )


def interpretar_pink_sheet(
    precios: list[tuple],
    descripcion: list[tuple],
    descripcion_oro: str = PINK_SHEET_DESCRIPCION_ORO,
    descripcion_plata: str = PINK_SHEET_DESCRIPCION_PLATA,
) -> PinkSheet:
    """Saca oro y plata de las filas de las hojas 'Monthly Prices' y 'Description'.

    Las descripciones esperadas son las de la edición vigente, salvo que se lea
    otra edición: cada una se verifica contra la frase que tenía cuando se leyó.
    """
    fila_nombres = next(
        (
            i
            for i, fila in enumerate(precios[:15])
            if any(isinstance(c, str) and c.strip() == "Gold" for c in fila)
        ),
        None,
    )
    if fila_nombres is None:
        raise ErrorDeFuente("Pink Sheet: no se encontró la columna 'Gold' en el encabezado")
    nombres = [c.strip() if isinstance(c, str) else c for c in precios[fila_nombres]]
    unidades = precios[fila_nombres + 1]

    columnas = {}
    for nombre in ("Gold", "Silver"):
        if nombre not in nombres:
            raise ErrorDeFuente(f"Pink Sheet: no se encontró la columna '{nombre}'")
        columna = nombres.index(nombre)
        unidad = unidades[columna]
        if not isinstance(unidad, str) or normalizar_texto(unidad) != normalizar_texto(PINK_SHEET_UNIDAD):
            raise ErrorDeFuente(
                f"Pink Sheet: la unidad de '{nombre}' es {unidad!r}, se esperaba "
                f"{PINK_SHEET_UNIDAD!r}"
            )
        columnas[nombre] = columna

    _verificar_descripcion(descripcion, descripcion_oro, "del oro")
    _verificar_descripcion(descripcion, descripcion_plata, "de la plata")

    oro: list[tuple[pd.Timestamp, float]] = []
    plata: list[tuple[pd.Timestamp, float]] = []
    for fila in precios[fila_nombres + 2 :]:
        etiqueta = fila[0].strip() if fila and isinstance(fila[0], str) else ""
        if not re.fullmatch(r"\d{4}M\d{2}", etiqueta):
            continue
        mes = pd.Timestamp(year=int(etiqueta[:4]), month=int(etiqueta[5:7]), day=1)
        for destino, nombre in ((oro, "Gold"), (plata, "Silver")):
            valor = fila[columnas[nombre]]
            # El Pink Sheet marca lo que no tiene con "…". Un hueco es un hueco.
            if isinstance(valor, (int, float)) and not isinstance(valor, bool):
                destino.append((mes, float(valor)))
    if not oro or not plata:
        raise ErrorDeFuente("Pink Sheet: no se pudo leer ninguna fila de oro o de plata")
    return PinkSheet(
        oro=_serie_mensual(oro, "oro"),
        plata=_serie_mensual(plata, "plata"),
        actualizada=_fecha_de_actualizacion(precios),
    )


def leer_pink_sheet(
    ruta: Path,
    descripcion_oro: str = PINK_SHEET_DESCRIPCION_ORO,
    descripcion_plata: str = PINK_SHEET_DESCRIPCION_PLATA,
) -> PinkSheet:
    return interpretar_pink_sheet(
        filas_de_xlsx(ruta, PINK_SHEET_HOJA_PRECIOS),
        filas_de_xlsx(ruta, PINK_SHEET_HOJA_DESCRIPCION),
        descripcion_oro,
        descripcion_plata,
    )


# --- FMI, Primary Commodity Prices -------------------------------------------


@dataclass(frozen=True)
class PreciosFMI:
    oro: pd.Series
    plata: pd.Series


def interpretar_fmi(filas: list[tuple]) -> PreciosFMI:
    """Saca oro y plata de las filas de la hoja 'External' de la base mensual.

    La primera fila trae el código de cada serie y la segunda, su descripción.
    La descripción es la convención: si cambia, la corrida se detiene.
    """
    if len(filas) < 3 or not isinstance(filas[0][0], str) or filas[0][0].strip() != "Commodity":
        raise ErrorDeFuente("FMI: la hoja no empieza con la fila de códigos ('Commodity')")
    codigos = [c.strip() if isinstance(c, str) else c for c in filas[0]]
    descripciones = filas[1]

    series = {}
    for nombre, codigo, esperada in (
        ("oro", FMI_CODIGO_ORO, FMI_DESCRIPCION_ORO),
        ("plata", FMI_CODIGO_PLATA, FMI_DESCRIPCION_PLATA),
    ):
        if codigo not in codigos:
            raise ErrorDeFuente(f"FMI: no se encontró la serie '{codigo}'")
        columna = codigos.index(codigo)
        hallada = descripciones[columna] if columna < len(descripciones) else None
        if not isinstance(hallada, str) or normalizar_texto(hallada) != normalizar_texto(esperada):
            raise ErrorDeFuente(
                f"FMI: la descripción de {codigo} ya no es la que se leyó al fijar el "
                f"control.\n  esperada: {esperada}\n  hallada:  {hallada}\n"
                "Releer la fuente y actualizar configuracion.py, FUENTES.md y "
                "SUPUESTOS.md (A-R0-20)."
            )
        pares: list[tuple[pd.Timestamp, float]] = []
        for fila in filas[2:]:
            etiqueta = fila[0].strip() if fila and isinstance(fila[0], str) else ""
            # El FMI escribe los meses sin cero a la izquierda: 1980M1, 1980M10.
            hallado = re.fullmatch(r"(\d{4})M(\d{1,2})", etiqueta)
            if not hallado:
                continue
            valor = fila[columna] if columna < len(fila) else None
            if isinstance(valor, (int, float)) and not isinstance(valor, bool):
                mes = pd.Timestamp(year=int(hallado.group(1)), month=int(hallado.group(2)), day=1)
                pares.append((mes, float(valor)))
        if not pares:
            raise ErrorDeFuente(f"FMI: no se pudo leer ninguna fila de {codigo}")
        series[nombre] = _serie_mensual(pares, f"fmi_{nombre}")
    return PreciosFMI(oro=series["oro"], plata=series["plata"])


def leer_fmi(ruta: Path) -> PreciosFMI:
    return interpretar_fmi(filas_de_xlsx(ruta, FMI_HOJA))


# --- Shiller -----------------------------------------------------------------


def filas_de_xls(ruta: Path, hoja: str) -> list[tuple]:
    import xlrd

    libro = xlrd.open_workbook(str(ruta))
    if hoja not in libro.sheet_names():
        raise ErrorDeFuente(f"{ruta.name}: no tiene la hoja '{hoja}' (tiene {libro.sheet_names()})")
    tabla = libro.sheet_by_name(hoja)
    return [tuple(tabla.row_values(i)) for i in range(tabla.nrows)]


def interpretar_shiller(filas: list[tuple]) -> pd.Series:
    """Saca la columna P (S&P Comp.) de la hoja 'Data' de ie_data.xls.

    La fecha viene como un número, 2026.09 o 1871.1 (octubre de 1871): la parte
    entera es el año y los dos primeros decimales, el mes.
    """
    encabezado = next(
        (
            i
            for i, fila in enumerate(filas[:30])
            if len(fila) > 1 and str(fila[0]).strip() == "Date" and str(fila[1]).strip() == "P"
        ),
        None,
    )
    if encabezado is None:
        raise ErrorDeFuente("Shiller: no se encontró el encabezado 'Date' / 'P' en la hoja Data")

    pares: list[tuple[pd.Timestamp, float]] = []
    for fila in filas[encabezado + 1 :]:
        if len(fila) < 2 or not isinstance(fila[0], float) or not 1800 < fila[0] < 2200:
            continue
        if not isinstance(fila[1], (int, float)):
            continue  # fila sin precio
        anio = int(fila[0])
        mes = round((fila[0] - anio) * 100)
        if not 1 <= mes <= 12:
            raise ErrorDeFuente(f"Shiller: la fecha {fila[0]!r} no es un año.mes")
        pares.append((pd.Timestamp(year=anio, month=mes, day=1), float(fila[1])))
    if not pares:
        raise ErrorDeFuente("Shiller: no se pudo leer ninguna fila de precios")
    return _serie_mensual(pares, "sp500")


def leer_shiller(ruta: Path) -> pd.Series:
    return interpretar_shiller(filas_de_xls(ruta, "Data"))


# --- FRED, series diarias ----------------------------------------------------


def interpretar_fred_diario(texto: str, id_serie: str) -> pd.Series:
    """Lee un CSV de fredgraph. Los días sin dato vienen vacíos o con un punto."""
    tabla = pd.read_csv(StringIO(texto))
    if tabla.shape[1] < 2 or normalizar_texto(str(tabla.columns[1])) != normalizar_texto(id_serie):
        raise ErrorDeFuente(
            f"{id_serie}: el CSV de FRED trae las columnas {list(tabla.columns)}, "
            f"se esperaba la fecha y '{id_serie}'"
        )
    fechas = pd.to_datetime(tabla.iloc[:, 0], errors="coerce")
    valores = pd.to_numeric(tabla.iloc[:, 1], errors="coerce")
    validas = fechas.notna()
    return _serie_diaria(fechas[validas], valores[validas], id_serie)


def leer_fred_diario(ruta: Path, id_serie: str) -> pd.Series:
    return interpretar_fred_diario(ruta.read_text(encoding="utf-8"), id_serie)


# --- Coin Metrics ------------------------------------------------------------


def interpretar_coin_metrics(documento: dict, metrica: str = "PriceUSD") -> pd.Series:
    """Lee la respuesta de asset-metrics. La fila del día D es el fixing de las
    00:00 UTC del día D+1, es decir, el cierre del día D."""
    if documento.get("next_page_url"):
        raise ErrorDeFuente(
            "Coin Metrics: la historia ya no entra en un solo pedido (la respuesta trae "
            "next_page_url). Leer una sola página truncaría la serie en silencio; hay que "
            "implementar la paginación antes de seguir."
        )
    filas = documento.get("data")
    if not isinstance(filas, list) or not filas:
        raise ErrorDeFuente("Coin Metrics: la respuesta no trae filas en 'data'")
    fechas, valores = [], []
    for fila in filas:
        if fila.get(metrica) in (None, ""):
            continue
        fechas.append(str(fila["time"])[:10])
        valores.append(float(fila[metrica]))
    return _serie_diaria(fechas, valores, "btc")


def leer_coin_metrics(ruta: Path) -> pd.Series:
    try:
        documento = json.loads(ruta.read_text(encoding="utf-8"))
    except json.JSONDecodeError as error:
        raise ErrorDeFuente(f"Coin Metrics: {ruta.name} no es JSON ({error})") from error
    return interpretar_coin_metrics(documento)


# --- Contrastes: se leen y no se guardan (A-R0-12) ---------------------------


def contraste_fred_diario(contraste: Contraste, id_serie: str, sesion=None) -> pd.Series:
    return interpretar_fred_diario(_pedir(contraste.url, sesion).text, id_serie)


def interpretar_nasdaq_api(documento: dict) -> pd.Series:
    try:
        filas = documento["data"]["tradesTable"]["rows"]
    except (KeyError, TypeError) as error:
        raise ErrorDeFuente(
            f"API de nasdaq.com: la respuesta no trae data.tradesTable.rows ({documento.get('status')})"
        ) from error
    fechas = [datetime.strptime(fila["date"], "%m/%d/%Y") for fila in filas]
    valores = [float(str(fila["close"]).replace(",", "")) for fila in filas]
    return _serie_diaria(fechas, valores, "nasdaq")


def contraste_nasdaq(contraste: Contraste, desde: date, hasta: date, sesion=None) -> pd.Series:
    respuesta = _pedir(
        contraste.url,
        sesion,
        cabeceras=CABECERAS_NAVEGADOR,
        parametros={
            "assetclass": "index",
            "fromdate": desde.isoformat(),
            "todate": hasta.isoformat(),
            "limit": 20000,
        },
    )
    try:
        return interpretar_nasdaq_api(respuesta.json())
    except ValueError as error:
        raise ErrorDeFuente(f"API de nasdaq.com: la respuesta no es JSON ({error})") from error


def interpretar_bitstamp(velas: list[dict]) -> pd.Series:
    fechas = [
        datetime.fromtimestamp(int(vela["timestamp"]), timezone.utc).strftime("%Y-%m-%d")
        for vela in velas
    ]
    return _serie_diaria(fechas, [float(vela["close"]) for vela in velas], "btc")


def contraste_bitstamp(contraste: Contraste, desde: date, hasta: date, sesion=None) -> pd.Series:
    """Velas diarias de Bitstamp, en ventanas de 1000 días, sin la vela abierta."""
    velas: list[dict] = []
    inicio = datetime(desde.year, desde.month, desde.day, tzinfo=timezone.utc)
    fin = datetime(hasta.year, hasta.month, hasta.day, tzinfo=timezone.utc)
    while inicio <= fin:
        respuesta = _pedir(
            contraste.url,
            sesion,
            parametros={
                "step": 86400,
                "limit": 1000,
                "start": int(inicio.timestamp()),
                "exclude_current_candle": "true",
            },
        )
        try:
            velas.extend(respuesta.json()["data"]["ohlc"])
        except (ValueError, KeyError, TypeError) as error:
            raise ErrorDeFuente(f"Bitstamp: la respuesta no trae data.ohlc ({error})") from error
        inicio += timedelta(days=1000)
    if not velas:
        raise ErrorDeFuente("Bitstamp: no devolvió ninguna vela")
    return interpretar_bitstamp(velas)
